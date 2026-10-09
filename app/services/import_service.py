from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from typing import Any, Literal, get_args

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session

from app.core.exceptions import APIError
from app.core.logging_utils import log_database_exception
from app.database.connection import atomic_transaction, rollback_failed_transaction
from app.models import (
    Call,
    Company,
    Contact,
    Customer,
    FollowUp,
    Meeting,
    SalesEnquiry,
)
from app.schemas.call import CallStatus
from app.schemas.customer import CustomerStatus, SalesStage
from app.schemas.follow_up import FollowUpStatus
from app.schemas.meeting import MeetingCreate, MeetingStatus
from app.schemas.sales_enquiry import EnquiryPriority, EnquiryStatus
from app.schemas.validators import (
    normalize_utc_datetime,
    resolve_local_datetime,
    timezone_name_from_datetime,
    timezone_from_name,
    validate_timezone_name,
)
from app.services.scheduling_service import (
    ScheduleRequest,
    default_duration,
    validate_schedule_batch,
)

import logging

logger = logging.getLogger(__name__)


EXPECTED_SHEETS = {
    "Companies",
    "Customers",
    "Contacts",
    "Sales_Enquiries",
    "Meetings",
    "Calls",
    "Follow_Ups",
    "TEST_EXPECTATIONS",
}


def _rows(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        return []

    headers = [
        str(value).strip() if value is not None else ""
        for value in rows[0]
    ]

    return [
        {
            headers[index]: value
            for index, value in enumerate(row)
            if index < len(headers) and headers[index]
        }
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


def _require_columns(
    rows: list[dict[str, Any]],
    sheet_name: str,
    required: set[str],
) -> None:
    if not rows:
        raise APIError(
            f"The {sheet_name} sheet is empty.",
            status_code=422,
            code="invalid_import",
        )

    actual = set(rows[0].keys())
    missing = required - actual

    if missing:
        raise APIError(
            f"{sheet_name} is missing columns: "
            + ", ".join(sorted(missing)),
            status_code=422,
            code="invalid_import",
        )


def _validate_enum_values(
    rows: list[dict[str, Any]],
    sheet_name: str,
    record_id: str,
    field: str,
    allowed_values: tuple[str, ...],
) -> None:
    for row in rows:
        if row[field] not in allowed_values:
            raise APIError(
                f"{sheet_name} record {row[record_id]} has an invalid {field}. "
                f"Expected one of: {', '.join(allowed_values)}.",
                status_code=422,
                code="invalid_import",
            )


def _datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, str):
        try:
            return datetime.fromisoformat(
                value.strip().replace("Z", "+00:00")
            )
        except ValueError as exc:
            raise APIError(
                f"Invalid datetime value: {value}",
                status_code=422,
                code="invalid_import",
            ) from exc

    raise APIError(
        f"Invalid datetime value: {value!r}",
        status_code=422,
        code="invalid_import",
    )


def _decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None

    try:
        return Decimal(str(value))
    except Exception as exc:
        raise APIError(
            f"Invalid monetary value: {value!r}",
            status_code=422,
            code="invalid_import",
        ) from exc


def _optional(row: dict[str, Any], key: str) -> Any:
    value = row.get(key)
    return None if value == "" else value


def _schedule_duration(
    row: dict[str, Any],
    *,
    sheet: str,
    record_id: Any,
    activity_type: Literal["meeting", "call", "followup"],
) -> int:
    value = row.get("duration")
    if value is None or value == "":
        return default_duration(activity_type)
    try:
        parsed = Decimal(str(value))
        if not parsed.is_finite() or parsed != parsed.to_integral_value():
            raise ValueError
        duration = int(parsed)
    except (ArithmeticError, TypeError, ValueError) as exc:
        raise APIError(
            f"{sheet} record {record_id} has an invalid duration.",
            status_code=422,
            code="invalid_import",
        ) from exc
    if not 1 <= duration <= 1440:
        raise APIError(
            f"{sheet} record {record_id} has an invalid duration.",
            status_code=422,
            code="invalid_import",
        )
    return duration


def _schedule_instant(
    value: datetime,
    *,
    timezone_name: str,
    occurrence: Any,
    sheet: str,
    record_id: Any,
) -> tuple[datetime, str]:
    try:
        timezone_name = validate_timezone_name(str(timezone_name))
        if value.tzinfo is None or value.utcoffset() is None:
            value = resolve_local_datetime(
                value.date(),
                value.time(),
                timezone_name,
                occurrence if occurrence in {"earlier", "later"} else None,
            )
        else:
            timezone_from_name(timezone_name)
        instant = normalize_utc_datetime(value)
        assert instant is not None
        return instant, timezone_name
    except (TypeError, ValueError) as exc:
        raise APIError(
            f"{sheet} record {record_id} has an invalid local scheduled time "
            f"for timezone {timezone_name}.",
            status_code=422,
            code="invalid_import",
        ) from exc


def _validate_relationships(
    rows: list[dict[str, Any]],
    *,
    sheet_name: str,
    record_id_field: str,
    customer_ids: set[Any],
    relationships: dict[str, tuple[str, dict[Any, Any]]],
) -> None:
    for row_number, row in enumerate(rows, start=2):
        customer_id = row["customer_id"]
        record_id = row[record_id_field]
        if customer_id not in customer_ids:
            raise APIError(
                f"{sheet_name} row {row_number} (record {record_id}) references "
                f"missing customer {customer_id}.",
                status_code=422,
                code="invalid_import",
            )

        for field, (label, customer_by_id) in relationships.items():
            related_id = _optional(row, field)
            if related_id is None:
                continue
            related_customer_id = customer_by_id.get(related_id)
            if related_customer_id is None:
                raise APIError(
                    f"{sheet_name} row {row_number} (record {record_id}) "
                    f"references missing {label} {related_id}.",
                    status_code=422,
                    code="invalid_import",
                )
            if related_customer_id != customer_id:
                raise APIError(
                    f"{sheet_name} row {row_number} (record {record_id}) "
                    f"references {label} {related_id} belonging to customer "
                    f"{related_customer_id}, not customer {customer_id}.",
                    status_code=422,
                    code="invalid_import",
                )


def import_crm_workbook(
    db: Session,
    file_bytes: bytes,
) -> dict[str, int]:
    try:
        workbook = load_workbook(
            filename=BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise APIError(
            "The uploaded file is not a valid Excel workbook.",
            status_code=422,
            code="invalid_import",
        ) from exc

    actual_sheets = set(workbook.sheetnames)
    missing_sheets = EXPECTED_SHEETS - actual_sheets

    if missing_sheets:
        raise APIError(
            "Workbook is missing sheets: "
            + ", ".join(sorted(missing_sheets)),
            status_code=422,
            code="invalid_import",
        )

    companies = _rows(workbook, "Companies")
    customers = _rows(workbook, "Customers")
    contacts = _rows(workbook, "Contacts")
    enquiries = _rows(workbook, "Sales_Enquiries")
    meetings = _rows(workbook, "Meetings")
    calls = _rows(workbook, "Calls")
    followups = _rows(workbook, "Follow_Ups")

    _require_columns(
        companies,
        "Companies",
        {
            "company_id",
            "company_name",
            "industry",
            "website",
            "address",
            "city",
            "country",
            "company_size",
            "description",
        },
    )

    _require_columns(
        customers,
        "Customers",
        {
            "customer_id",
            "company_id",
            "customer_name",
            "status",
            "sales_stage",
        },
    )

    _require_columns(
        contacts,
        "Contacts",
        {
            "contact_id",
            "customer_id",
            "name",
            "job_title",
            "email",
            "phone",
            "is_primary",
        },
    )

    _require_columns(
        enquiries,
        "Sales_Enquiries",
        {
            "enquiry_id",
            "customer_id",
            "product",
            "enquiry_text",
            "priority",
            "status",
            "estimated_value",
            "created_at",
        },
    )

    _require_columns(
        meetings,
        "Meetings",
        {
            "meeting_id",
            "customer_id",
            "contact_id",
            "enquiry_id",
            "scheduled_at",
            "duration",
            "status",
            "agenda",
            "notes",
            "summary",
        },
    )

    _require_columns(
        calls,
        "Calls",
        {
            "call_id",
            "customer_id",
            "contact_id",
            "enquiry_id",
            "call_type",
            "scheduled_at",
            "actual_time",
            "status",
            "outcome",
            "notes",
            "summary",
            "next_followup_date",
        },
    )

    _require_columns(
        followups,
        "Follow_Ups",
        {
            "followup_id",
            "customer_id",
            "enquiry_id",
            "meeting_id",
            "call_id",
            "type",
            "due_date",
            "status",
            "description",
            "assigned_to",
            "completed_at",
        },
    )

    for rows, sheet, record_id, field, values in (
        (customers, "Customers", "customer_id", "status", get_args(CustomerStatus)),
        (customers, "Customers", "customer_id", "sales_stage", get_args(SalesStage)),
        (enquiries, "Sales_Enquiries", "enquiry_id", "priority", get_args(EnquiryPriority)),
        (enquiries, "Sales_Enquiries", "enquiry_id", "status", get_args(EnquiryStatus)),
        (meetings, "Meetings", "meeting_id", "status", get_args(MeetingStatus)),
        (calls, "Calls", "call_id", "status", get_args(CallStatus)),
        (followups, "Follow_Ups", "followup_id", "status", get_args(FollowUpStatus)),
    ):
        _validate_enum_values(rows, sheet, record_id, field, values)

    meeting_schedules: list[tuple[datetime, str, int]] = []
    for row in meetings:
        try:
            meeting_input = MeetingCreate.model_validate(
                {
                    "customer_id": row["customer_id"],
                    "scheduled_at": _datetime(row.get("scheduled_at")),
                    "scheduled_timezone": _optional(row, "scheduled_timezone"),
                }
            )
        except (ValidationError, TypeError, ValueError) as exc:
            raise APIError(
                f"Meeting {row['meeting_id']} must have a valid timezone-aware "
                "scheduled_at timestamp.",
                status_code=422,
                code="invalid_import",
            ) from exc
        scheduled_at = normalize_utc_datetime(meeting_input.scheduled_at)
        assert scheduled_at is not None
        meeting_schedules.append(
            (
                scheduled_at,
                meeting_input.scheduled_timezone or "UTC",
                _schedule_duration(
                    row,
                    sheet="Meetings",
                    record_id=row["meeting_id"],
                    activity_type="meeting",
                ),
            )
        )

    call_schedules: dict[Any, tuple[datetime | None, str, int]] = {}
    followup_schedules: dict[Any, tuple[datetime, str, int]] = {}
    for row in calls:
        scheduled_at = _datetime(row.get("scheduled_at"))
        if scheduled_at is not None:
            timezone_name = _optional(row, "scheduled_timezone")
            if timezone_name is None:
                timezone_name = (
                    timezone_name_from_datetime(scheduled_at)
                    if scheduled_at.tzinfo is not None
                    and scheduled_at.utcoffset() is not None
                    else "UTC"
                )
            scheduled_at, timezone_name = _schedule_instant(
                scheduled_at,
                timezone_name=timezone_name,
                occurrence=_optional(row, "scheduled_time_occurrence"),
                sheet="Calls",
                record_id=row["call_id"],
            )
        else:
            timezone_name = "UTC"
        duration = _schedule_duration(
            row,
            sheet="Calls",
            record_id=row["call_id"],
            activity_type="call",
        )
        call_schedules[row["call_id"]] = (scheduled_at, timezone_name, duration)

    for row in followups:
        due_date = _datetime(row.get("due_date"))
        if due_date is None:
            raise APIError(
                f"Follow_Ups record {row['followup_id']} must have a due_date.",
                status_code=422,
                code="invalid_import",
            )
        timezone_name = _optional(row, "due_timezone")
        if timezone_name is None:
            timezone_name = (
                timezone_name_from_datetime(due_date)
                if due_date.tzinfo is not None and due_date.utcoffset() is not None
                else "UTC"
            )
        due_date, timezone_name = _schedule_instant(
            due_date,
            timezone_name=timezone_name,
            occurrence=_optional(row, "due_time_occurrence"),
            sheet="Follow_Ups",
            record_id=row["followup_id"],
        )
        duration = _schedule_duration(
            row,
            sheet="Follow_Ups",
            record_id=row["followup_id"],
            activity_type="followup",
        )
        followup_schedules[row["followup_id"]] = (
            due_date,
            timezone_name,
            duration,
        )

    # ---------------------------------------------------------
    # Validate all relationships before inserting anything.
    # ---------------------------------------------------------

    company_ids = {row["company_id"] for row in companies}
    customer_ids = {row["customer_id"] for row in customers}
    contact_customer_ids = {
        row["contact_id"]: row["customer_id"] for row in contacts
    }
    enquiry_customer_ids = {
        row["enquiry_id"]: row["customer_id"] for row in enquiries
    }
    meeting_customer_ids = {
        row["meeting_id"]: row["customer_id"] for row in meetings
    }
    call_customer_ids = {
        row["call_id"]: row["customer_id"] for row in calls
    }

    for row in customers:
        if row["company_id"] not in company_ids:
            raise APIError(
                f"Customer {row['customer_id']} references "
                f"missing company {row['company_id']}.",
                status_code=422,
                code="invalid_import",
            )

    _validate_relationships(
        contacts,
        sheet_name="Contacts",
        record_id_field="contact_id",
        customer_ids=customer_ids,
        relationships={},
    )
    _validate_relationships(
        enquiries,
        sheet_name="Sales_Enquiries",
        record_id_field="enquiry_id",
        customer_ids=customer_ids,
        relationships={},
    )
    _validate_relationships(
        meetings,
        sheet_name="Meetings",
        record_id_field="meeting_id",
        customer_ids=customer_ids,
        relationships={
            "contact_id": ("contact", contact_customer_ids),
            "enquiry_id": ("enquiry", enquiry_customer_ids),
        },
    )
    _validate_relationships(
        calls,
        sheet_name="Calls",
        record_id_field="call_id",
        customer_ids=customer_ids,
        relationships={
            "contact_id": ("contact", contact_customer_ids),
            "enquiry_id": ("enquiry", enquiry_customer_ids),
        },
    )
    _validate_relationships(
        followups,
        sheet_name="Follow_Ups",
        record_id_field="followup_id",
        customer_ids=customer_ids,
        relationships={
            "enquiry_id": ("enquiry", enquiry_customer_ids),
            "meeting_id": ("meeting", meeting_customer_ids),
            "call_id": ("call", call_customer_ids),
        },
    )

    # ---------------------------------------------------------
    # Transaction
    #
    # If ANY insert fails, the entire import is rolled back.
    # ---------------------------------------------------------

    try:
        with atomic_transaction(db):
            schedule_requests: list[ScheduleRequest] = []
            for row, (scheduled_at, timezone_name, duration) in zip(
                meetings, meeting_schedules
            ):
                if row["status"] == "scheduled":
                    schedule_requests.append(
                        ScheduleRequest(
                            "meeting",
                            scheduled_at,
                            duration,
                            timezone_name,
                            record_id=row["meeting_id"],
                        )
                    )
            for row in calls:
                scheduled_at, timezone_name, duration = call_schedules[
                    row["call_id"]
                ]
                if row["status"] == "scheduled" and scheduled_at is not None:
                    schedule_requests.append(
                        ScheduleRequest(
                            "call",
                            scheduled_at,
                            duration,
                            timezone_name,
                            record_id=row["call_id"],
                        )
                    )
            for row in followups:
                due_date, timezone_name, duration = followup_schedules[
                    row["followup_id"]
                ]
                if row["status"] in {"pending", "in_progress", "overdue"}:
                    schedule_requests.append(
                        ScheduleRequest(
                            "followup",
                            due_date,
                            duration,
                            timezone_name,
                            record_id=row["followup_id"],
                        )
                    )
            validate_schedule_batch(db, schedule_requests)

            for row in companies:
                db.add(
                    Company(
                        company_id=row["company_id"],
                        company_name=row["company_name"],
                        industry=_optional(row, "industry"),
                        website=_optional(row, "website"),
                        address=_optional(row, "address"),
                        city=_optional(row, "city"),
                        country=_optional(row, "country"),
                        company_size=_optional(row, "company_size"),
                        description=_optional(row, "description"),
                    )
                )

            db.flush()

            for row in customers:
                db.add(
                    Customer(
                        customer_id=row["customer_id"],
                        company_id=row["company_id"],
                        customer_name=row["customer_name"],
                        status=row["status"],
                        sales_stage=row["sales_stage"],
                    )
                )

            db.flush()

            for row in contacts:
                db.add(
                    Contact(
                        contact_id=row["contact_id"],
                        customer_id=row["customer_id"],
                        name=row["name"],
                        job_title=_optional(row, "job_title"),
                        email=_optional(row, "email"),
                        phone=_optional(row, "phone"),
                        is_primary=bool(row.get("is_primary", False)),
                    )
                )

            db.flush()

            for row in enquiries:
                db.add(
                    SalesEnquiry(
                        enquiry_id=row["enquiry_id"],
                        customer_id=row["customer_id"],
                        product=_optional(row, "product"),
                        enquiry_text=row["enquiry_text"],
                        priority=row["priority"],
                        status=row["status"],
                        estimated_value=_decimal(
                            row.get("estimated_value")
                        ),
                        created_at=_datetime(row["created_at"]),
                    )
                )

            db.flush()

            for row, (scheduled_at, scheduled_timezone, duration) in zip(
                meetings, meeting_schedules
            ):
                db.add(
                    Meeting(
                        meeting_id=row["meeting_id"],
                        customer_id=row["customer_id"],
                        contact_id=_optional(row, "contact_id"),
                        enquiry_id=_optional(row, "enquiry_id"),
                        scheduled_at=scheduled_at,
                        scheduled_timezone=scheduled_timezone,
                        duration=duration,
                        status=row["status"],
                        agenda=_optional(row, "agenda"),
                        notes=_optional(row, "notes"),
                        summary=_optional(row, "summary"),
                    )
                )

            db.flush()

            for row in calls:
                scheduled_at, scheduled_timezone, duration = call_schedules[
                    row["call_id"]
                ]
                db.add(
                    Call(
                        call_id=row["call_id"],
                        customer_id=row["customer_id"],
                        contact_id=_optional(row, "contact_id"),
                        enquiry_id=_optional(row, "enquiry_id"),
                        call_type=_optional(row, "call_type"),
                        scheduled_at=scheduled_at,
                        scheduled_timezone=scheduled_timezone,
                        duration=duration if scheduled_at is not None else _optional(row, "duration"),
                        actual_time=_datetime(
                            row.get("actual_time")
                        ),
                        status=row["status"],
                        outcome=_optional(row, "outcome"),
                        notes=_optional(row, "notes"),
                        summary=_optional(row, "summary"),
                        next_followup_date=_datetime(
                            row.get("next_followup_date")
                        ),
                    )
                )

            db.flush()

            for row in followups:
                due_date, due_timezone, duration = followup_schedules[
                    row["followup_id"]
                ]
                db.add(
                    FollowUp(
                        followup_id=row["followup_id"],
                        customer_id=row["customer_id"],
                        enquiry_id=_optional(row, "enquiry_id"),
                        meeting_id=_optional(row, "meeting_id"),
                        call_id=_optional(row, "call_id"),
                        type=row["type"],
                        due_date=due_date,
                        due_timezone=due_timezone,
                        duration=duration,
                        status=row["status"],
                        description=_optional(row, "description"),
                        assigned_to=_optional(row, "assigned_to"),
                        completed_at=_datetime(
                            row.get("completed_at")
                        ),
                    )
                )

            db.flush()

    except APIError:
        raise
    except IntegrityError as exc:
        rollback_failed_transaction(db)
        log_database_exception(
            logger,
            "import_crm_workbook",
            exc,
            conflict=True,
        )
        raise APIError(
            "The CRM import conflicts with existing database data.",
            status_code=409,
            code="import_conflict",
        ) from exc
    except SQLAlchemyError as exc:
        rollback_failed_transaction(db)
        log_database_exception(
            logger,
            "import_crm_workbook",
            exc,
        )
        raise APIError(
            "The CRM import could not be completed.",
            status_code=500,
            code="database_error",
        ) from exc

    return {
        "companies": len(companies),
        "customers": len(customers),
        "contacts": len(contacts),
        "sales_enquiries": len(enquiries),
        "meetings": len(meetings),
        "calls": len(calls),
        "follow_ups": len(followups),
    }