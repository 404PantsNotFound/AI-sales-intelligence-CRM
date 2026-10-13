"""Organization isolation: a member of A can never read, change or use B's records.

Two organizations hold overlapping data (the same company name, the same meeting
times, and B's rows created first so they own the low IDs). Each test proves a
boundary at the route and, where the route is not the only way in, at the service.
"""

from collections.abc import Generator
from dataclasses import dataclass
from datetime import datetime, timezone
import json
from typing import Any

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.orm import Session, sessionmaker

from app.agent.tools import AgentToolContext, build_read_only_tools
from app.core.config import settings
from app.core.exceptions import APIError
from app.core.security import create_access_token
from app.models import (
    Call,
    CalendarEventLink,
    Company,
    Contact,
    Customer,
    FollowUp,
    GoogleCalendarConnection,
    HitlActionProposal,
    Meeting,
    Organization,
    SalesEnquiry,
    SchedulingLock,
    User,
)
from app.models.tenancy import CrossOrganizationReference
from app.services import calendar_sync_service, customer_service
from app.services.agent_action_state import pending_actions
from app.services.hitl_policy import get_action_policy
from app.services.scheduling_service import ScheduleRequest, inspect_availability
from tests.conftest import TEST_USER_PASSWORD_HASH, legacy_organization_id
from tests.test_customers_api import imported_workbook
from tests.test_google_calendar import FakeGoogle, google  # noqa: F401  (fixture)

SLOT = datetime(2026, 10, 5, 10, 0, tzinfo=timezone.utc)  # A Monday after the frozen "now".
SHARED_COMPANY = "Shared Name Ltd"


@dataclass
class Tenant:
    organization_id: int
    user_id: int
    admin_id: int
    company_id: int
    customer_id: int
    contact_id: int
    enquiry_id: int
    meeting_id: int
    call_id: int
    followup_id: int


def _seed_tenant(db: Session, organization_id: int, label: str) -> Tenant:
    users = [
        User(
            email=f"{label}.{role}@example.com",
            full_name=f"{label} {role}",
            password_hash=TEST_USER_PASSWORD_HASH,
            role=role,
            organization_id=organization_id,
        )
        for role in ("sales", "admin")
    ]
    company = Company(organization_id=organization_id, company_name=SHARED_COMPANY)
    customer = Customer(
        company=company, customer_name=f"{label} Customer", status="active", sales_stage="qualified"
    )
    contact = Contact(customer=customer, name=f"{label} Contact", email=f"{label}@example.com")
    enquiry = SalesEnquiry(
        customer=customer, product=f"{label} product", enquiry_text="Needs a CRM",
        priority="high", status="open", estimated_value=1000,
        created_at=datetime(2026, 8, 1, tzinfo=timezone.utc),
    )
    meeting = Meeting(
        customer=customer, contact=contact, enquiry=enquiry, scheduled_at=SLOT,
        duration=60, status="scheduled", agenda=f"{label} secret agenda",
    )
    call = Call(
        customer=customer, contact=contact, enquiry=enquiry, call_type="Discovery",
        scheduled_at=datetime(2026, 10, 6, 10, 0, tzinfo=timezone.utc), status="scheduled",
    )
    followup = FollowUp(
        customer=customer, enquiry=enquiry, meeting=meeting, call=call, type="email",
        due_date=datetime(2026, 10, 7, 10, 0, tzinfo=timezone.utc), status="pending",
        description=f"{label} follow-up",
    )
    db.add_all([*users, company, customer, contact, enquiry, meeting, call, followup])
    db.flush()
    return Tenant(
        organization_id, users[0].user_id, users[1].user_id, company.company_id,
        customer.customer_id, contact.contact_id, enquiry.enquiry_id, meeting.meeting_id,
        call.call_id, followup.followup_id,
    )


@pytest.fixture
def tenants(agent_sessions: sessionmaker[Session]) -> tuple[Tenant, Tenant]:
    with agent_sessions() as db:
        a_org = legacy_organization_id(db)
        rival = Organization(name="Rival Corp", slug="rival-corp")
        db.add(rival)
        db.flush()
        db.add(SchedulingLock(lock_id=rival.organization_id))
        # B first, so B owns the low IDs an attacker would guess.
        b = _seed_tenant(db, rival.organization_id, "bravo")
        a = _seed_tenant(db, a_org, "alpha")
        db.commit()
    return a, b


@pytest.fixture
def client(agent_client: TestClient, tenants: tuple[Tenant, Tenant]) -> TestClient:
    return _as(agent_client, tenants[0].user_id)


def _as(client: TestClient, user_id: int) -> TestClient:
    user = type("U", (), {"user_id": user_id, "email": "", "role": ""})()
    token, _ = create_access_token(user)  # type: ignore[arg-type]
    client.headers["Authorization"] = f"Bearer {token}"
    return client


def _not_found(response: Any) -> None:
    assert response.status_code == 404, response.text
    assert "bravo" not in response.text.lower()


# --- reads ------------------------------------------------------------------------


def test_lists_and_searches_show_only_own_organization(client: TestClient, tenants) -> None:
    a, b = tenants
    listed = client.get("/api/customers").json()
    assert [item["customer_id"] for item in listed["items"]] == [a.customer_id]
    assert listed["total"] == 1
    searched = client.get("/api/customers", params={"search": SHARED_COMPANY}).json()
    assert [item["customer_id"] for item in searched["items"]] == [a.customer_id]


@pytest.mark.parametrize(
    "path",
    [
        "/api/customers/{customer_id}",
        "/api/customers/{customer_id}/overview",
        "/api/customers/{customer_id}/activity",
        "/api/customers/{customer_id}/contacts",
        "/api/customers/{customer_id}/enquiries",
        "/api/customers/{customer_id}/meetings",
        "/api/customers/{customer_id}/calls",
        "/api/customers/{customer_id}/followups",
        "/api/contacts/{contact_id}",
        "/api/enquiries/{enquiry_id}",
        "/api/meetings/{meeting_id}",
        "/api/calls/{call_id}",
        "/api/followups/{followup_id}",
    ],
)
def test_foreign_records_look_missing(client: TestClient, tenants, path: str) -> None:
    a, b = tenants
    assert client.get(path.format(**vars(a))).status_code == 200
    _not_found(client.get(path.format(**vars(b))))


def test_missing_and_foreign_ids_get_the_same_response(client: TestClient, tenants) -> None:
    _, b = tenants
    foreign = client.get(f"/api/customers/{b.customer_id}")
    missing = client.get("/api/customers/999999")
    assert foreign.status_code == missing.status_code == 404
    assert foreign.json() == missing.json()


# --- writes -----------------------------------------------------------------------


def test_updates_to_foreign_records_are_rejected_and_change_nothing(
    client: TestClient, tenants, agent_sessions: sessionmaker[Session]
) -> None:
    _, b = tenants
    _not_found(client.patch(f"/api/customers/{b.customer_id}", json={"status": "inactive"}))
    _not_found(client.patch(f"/api/contacts/{b.contact_id}", json={"name": "Hijacked"}))
    _not_found(client.patch(f"/api/enquiries/{b.enquiry_id}", json={"status": "lost"}))
    _not_found(client.put(f"/api/meetings/{b.meeting_id}", json={"status": "cancelled"}))
    _not_found(client.put(f"/api/calls/{b.call_id}", json={"status": "cancelled"}))
    _not_found(client.put(f"/api/followups/{b.followup_id}", json={"status": "cancelled"}))
    _not_found(
        client.post(
            f"/api/scheduling/activities/meeting/{b.meeting_id}/move",
            json={"target_date": "2026-10-08", "target_timezone": "UTC"},
        )
    )
    with agent_sessions() as db:
        assert db.get(Customer, b.customer_id).status == "active"
        assert db.get(Contact, b.contact_id).name == "bravo Contact"
        assert db.get(SalesEnquiry, b.enquiry_id).status == "open"
        assert db.get(Meeting, b.meeting_id).status == "scheduled"
        assert db.get(Meeting, b.meeting_id).scheduled_at.date().isoformat() == "2026-10-05"
        assert db.get(Call, b.call_id).status == "scheduled"
        assert db.get(FollowUp, b.followup_id).status == "pending"


def test_creates_cannot_reference_foreign_records(
    client: TestClient, tenants, agent_sessions: sessionmaker[Session]
) -> None:
    a, b = tenants
    future = "2026-10-09T10:00:00+00:00"
    cases = [
        ("/api/contacts", {"customer_id": b.customer_id, "name": "Spy"}),
        ("/api/enquiries", {"customer_id": b.customer_id, "enquiry_text": "Spy"}),
        ("/api/meetings", {"customer_id": b.customer_id, "scheduled_at": future}),
        ("/api/calls", {"customer_id": b.customer_id, "scheduled_at": future}),
        ("/api/followups", {"customer_id": b.customer_id, "type": "email", "due_date": future}),
        # Own customer, foreign related record.
        ("/api/meetings", {"customer_id": a.customer_id, "contact_id": b.contact_id, "scheduled_at": future}),
        ("/api/calls", {"customer_id": a.customer_id, "enquiry_id": b.enquiry_id, "scheduled_at": future}),
        ("/api/followups", {"customer_id": a.customer_id, "meeting_id": b.meeting_id, "type": "email", "due_date": future}),
        ("/api/followups", {"customer_id": a.customer_id, "call_id": b.call_id, "type": "email", "due_date": future}),
    ]
    with agent_sessions() as db:
        before = [db.scalar(select(func.count()).select_from(m)) for m in (Contact, SalesEnquiry, Meeting, Call, FollowUp)]
    for path, body in cases:
        _not_found(client.post(path, json=body))
    with agent_sessions() as db:
        after = [db.scalar(select(func.count()).select_from(m)) for m in (Contact, SalesEnquiry, Meeting, Call, FollowUp)]
    assert after == before


def test_relationships_cannot_be_moved_to_foreign_records(client: TestClient, tenants) -> None:
    a, b = tenants
    _not_found(client.patch(f"/api/customers/{a.customer_id}", json={"company_id": b.company_id}))
    _not_found(client.patch(f"/api/contacts/{a.contact_id}", json={"customer_id": b.customer_id}))
    _not_found(client.patch(f"/api/enquiries/{a.enquiry_id}", json={"customer_id": b.customer_id}))
    _not_found(client.put(f"/api/meetings/{a.meeting_id}", json={"contact_id": b.contact_id}))
    _not_found(client.put(f"/api/calls/{a.call_id}", json={"customer_id": b.customer_id}))
    _not_found(client.put(f"/api/followups/{a.followup_id}", json={"call_id": b.call_id}))


def test_company_name_lookup_never_selects_another_organizations_company(
    client: TestClient, tenants, agent_sessions: sessionmaker[Session]
) -> None:
    a, b = tenants
    response = client.post(
        "/api/customers",
        json={
            "customer_name": "Second Buyer",
            "company": {"company_name": SHARED_COMPANY.upper()},
            "primary_contact": {"name": "New Contact"},
            "sales_enquiry": {"enquiry_text": "Needs more"},
        },
    )
    assert response.status_code == 201, response.text
    # The organization's own same-named company is reused; B's is never touched.
    assert response.json()["company"]["company_id"] == a.company_id
    with agent_sessions() as db:
        assert db.scalar(select(func.count()).select_from(Customer).where(Customer.company_id == b.company_id)) == 1


def test_orm_guard_rejects_cross_organization_links(
    agent_sessions: sessionmaker[Session], tenants
) -> None:
    a, b = tenants
    with agent_sessions() as db:
        db.add(Meeting(organization_id=a.organization_id, customer_id=a.customer_id,
                       contact_id=b.contact_id, scheduled_at=SLOT, status="completed"))
        with pytest.raises(CrossOrganizationReference):
            db.flush()
        db.rollback()
        db.add(Customer(organization_id=a.organization_id, company_id=b.company_id,
                        customer_name="Intruder", status="active", sales_stage="new"))
        with pytest.raises(CrossOrganizationReference):
            db.flush()
        db.rollback()
        # A child without an organization takes its parent's; it can never pick another.
        contact = Contact(customer_id=b.customer_id, name="Derived")
        db.add(contact)
        db.flush()
        assert contact.organization_id == b.organization_id
        db.rollback()
        existing = db.get(FollowUp, a.followup_id)
        existing.meeting_id = b.meeting_id
        with pytest.raises(CrossOrganizationReference):
            db.flush()


def test_service_lookups_require_the_callers_organization(
    agent_sessions: sessionmaker[Session], tenants
) -> None:
    a, b = tenants
    with agent_sessions() as db:
        with pytest.raises(APIError) as error:
            customer_service.get_customer(db, b.customer_id, organization_id=a.organization_id)
        assert error.value.status_code == 404
        customers, total = customer_service.list_customers(
            db, organization_id=b.organization_id, page=1, page_size=10
        )
        assert [c.customer_id for c in customers] == [b.customer_id] and total == 1


# --- analytics, dashboard and scheduling --------------------------------------------


def test_analytics_and_dashboard_exclude_foreign_data(client: TestClient, tenants) -> None:
    a, b = tenants
    overview = client.get("/api/analytics/overview").json()
    assert overview["total_customers"] == 1
    assert overview["open_enquiries"] == 1
    pipeline = client.get("/api/analytics/pipeline").json()
    assert sum(item["count"] for item in pipeline["by_sales_stage"]) == 1
    enquiries = client.get("/api/analytics/enquiries").json()
    assert float(enquiries["total_estimated_value"]) == 1000
    assert [item["label"] for item in enquiries["by_product"]] == ["alpha product"]
    attention = client.get("/api/analytics/attention").json()
    assert {item["customer_id"] for item in attention["items"]} <= {a.customer_id}
    activities = client.get(
        "/api/scheduling/activities",
        params={"start_at": "2026-10-01T00:00:00Z", "end_at": "2026-10-31T00:00:00Z"},
    ).json()["items"]
    assert activities and {item["customer_id"] for item in activities} == {a.customer_id}
    assert "bravo" not in json.dumps(activities)


def test_scheduling_conflicts_are_per_organization(
    client: TestClient, tenants, agent_sessions: sessionmaker[Session]
) -> None:
    a, b = tenants
    # B's meeting at SLOT is invisible to A: no conflict, no leaked times.
    with agent_sessions() as db:
        result = inspect_availability(
            db, ScheduleRequest("call", SLOT, 30, "UTC", a.organization_id)
        )
        own = inspect_availability(db, ScheduleRequest("call", SLOT, 30, "UTC", b.organization_id))
    # A's own meeting is also at SLOT, so A sees exactly one conflict: its own.
    assert result["available"] is False and len(result["conflicts"]) == 1
    assert own["available"] is False and len(own["conflicts"]) == 1
    response = client.post("/api/scheduling/availability", json={
        "activity_type": "call", "starts_at": "2026-10-06T10:00:00+00:00",
        "duration_minutes": 30, "timezone_name": "UTC",
    })
    # Both organizations have a call at this time; only A's own call conflicts.
    assert len(response.json()["conflicts"]) == 1


# --- agent ------------------------------------------------------------------------


def test_agent_read_tools_are_scoped_to_the_context_organization(
    agent_sessions: sessionmaker[Session], tenants
) -> None:
    a, b = tenants
    with agent_sessions() as db:
        tools = {t.name: t for t in build_read_only_tools(db, AgentToolContext(organization_id=a.organization_id))}
        found = json.loads(tools["find_customer"].invoke({"search": SHARED_COMPANY}))
        assert [item["customer_id"] for item in found["data"]["items"]] == [a.customer_id]
        for name in ("get_customer_profile", "get_sales_enquiries", "get_customer_meetings",
                     "get_customer_calls", "get_customer_followups", "get_customer_activity"):
            result = json.loads(tools[name].invoke({"customer_id": b.customer_id}))
            assert result["error"]["code"] == "customer_not_found", name
            assert "bravo" not in json.dumps(result)
        analytics = json.loads(tools["get_sales_analytics"].invoke({}))
        assert analytics["overview"]["total_customers"] == 1


def test_crm_tools_refuse_to_run_without_a_tenant_context() -> None:
    with pytest.raises(ValueError, match="tenant context"):
        build_read_only_tools(None)


def test_agent_routes_reject_foreign_customers_before_any_model_call(
    client: TestClient, tenants, monkeypatch: pytest.MonkeyPatch
) -> None:
    _, b = tenants

    def no_model() -> None:
        raise AssertionError("the model must not run for a foreign customer")

    monkeypatch.setattr("app.services.agent_service.create_chat_model", no_model)
    monkeypatch.setattr("app.services.ai_service.create_chat_model", no_model)
    _not_found(client.post("/api/agent/chat", json={"message": "Summarise", "customer_id": b.customer_id}))
    _not_found(client.post(f"/api/agent/customer-summary/{b.customer_id}"))
    _not_found(client.post(f"/api/agent/meeting-brief/{b.meeting_id}"))
    latest = client.get("/api/agent/conversations/latest", params={"customer_id": b.customer_id})
    assert latest.status_code == 200 and latest.json() is None


# --- HITL ---------------------------------------------------------------------------


def _pending(sessions: sessionmaker[Session], owner_id: int, payload: dict[str, int]) -> str:
    with sessions() as db:
        return pending_actions.add(db, "complete_followup", payload, "thread", user_id=owner_id).action_id


def test_tasks_cannot_target_foreign_records(client: TestClient, tenants) -> None:
    a, b = tenants
    response = client.post(
        "/api/agent/tasks",
        json={"operations": [{"action": "complete_followup",
                              "values": {"customer_id": b.customer_id, "followup_id": b.followup_id}}]},
    )
    _not_found(response)
    response = client.post(
        "/api/agent/tasks",
        json={"operations": [{"action": "complete_followup",
                              "values": {"customer_id": a.customer_id, "followup_id": b.followup_id}}]},
    )
    _not_found(response)


def test_task_form_choices_list_only_own_records(client: TestClient, tenants) -> None:
    a, _ = tenants
    response = client.post(
        "/api/agent/tasks", json={"operations": [{"action": "create_followup", "values": {}}]}
    )
    assert response.status_code == 200, response.text
    customer_field = next(
        field for field in response.json()["operations"][0]["fields"] if field["name"] == "customer_id"
    )
    assert [choice["id"] for choice in customer_field["choices"]] == [a.customer_id]


def test_admins_cannot_decide_another_organizations_actions(
    agent_client: TestClient, tenants, agent_sessions: sessionmaker[Session]
) -> None:
    a, b = tenants
    action_id = _pending(agent_sessions, b.user_id, {"customer_id": b.customer_id, "followup_id": b.followup_id})
    _as(agent_client, a.admin_id)
    for verb in ("confirm", "reject", "cancel"):
        _not_found(agent_client.post(f"/api/agent/actions/{action_id}/{verb}"))
    _not_found(agent_client.get(f"/api/agent/tasks/{_task_of(agent_sessions, action_id)}"))
    with agent_sessions() as db:
        assert db.get(HitlActionProposal, action_id).status == "pending"
        assert db.get(FollowUp, b.followup_id).status == "pending"
    # B's own admin still can.
    _as(agent_client, b.admin_id)
    assert agent_client.post(f"/api/agent/actions/{action_id}/confirm").json()["status"] == "completed"


def _task_of(sessions: sessionmaker[Session], action_id: str) -> str:
    with sessions() as db:
        return db.get(HitlActionProposal, action_id).task_id


def test_execution_revalidates_targets_against_the_owners_organization(
    client: TestClient, tenants, agent_sessions: sessionmaker[Session]
) -> None:
    a, b = tenants
    # A stored proposal naming B's records (as if proposal-time checks were bypassed).
    action_id = _pending(agent_sessions, a.user_id, {"customer_id": b.customer_id, "followup_id": b.followup_id})
    result = client.post(f"/api/agent/actions/{action_id}/confirm").json()
    assert result["status"] == "failed"
    assert result["error_code"] == "customer_not_found"
    with agent_sessions() as db:
        assert db.get(FollowUp, b.followup_id).status == "pending"


def test_pending_action_capacity_is_per_organization(
    tenants, agent_sessions: sessionmaker[Session], monkeypatch: pytest.MonkeyPatch
) -> None:
    a, b = tenants
    monkeypatch.setattr(settings, "pending_action_max_capacity", 1)
    _pending(agent_sessions, b.user_id, {"customer_id": b.customer_id, "followup_id": b.followup_id})
    # B is full, A is not.
    _pending(agent_sessions, a.user_id, {"customer_id": a.customer_id, "followup_id": a.followup_id})
    with pytest.raises(APIError) as error:
        _pending(agent_sessions, b.user_id, {"customer_id": b.customer_id, "followup_id": b.followup_id})
    assert error.value.code == "pending_actions_full"


def test_policy_overrides_stay_in_their_organization(
    agent_client: TestClient, tenants, agent_sessions: sessionmaker[Session]
) -> None:
    a, b = tenants
    _as(agent_client, a.admin_id)
    response = agent_client.put(
        "/api/agent/policies/create_followup", json={"action": "create_followup", "mode": "disabled"}
    )
    assert response.status_code == 200, response.text
    assert agent_client.get("/api/agent/policies").json()["create_followup"] == "disabled"
    _as(agent_client, b.admin_id)
    assert agent_client.get("/api/agent/policies").json()["create_followup"] == "approval_required"
    with agent_sessions() as db:
        assert get_action_policy(db, "create_followup", organization_id=b.organization_id) == "approval_required"
    # Ordinary users still cannot change policy, in any organization.
    _as(agent_client, b.user_id)
    forbidden = agent_client.put(
        "/api/agent/policies/create_followup", json={"action": "create_followup", "mode": "automatic"}
    )
    assert forbidden.status_code == 403


# --- imports and calendar -------------------------------------------------------------


def test_import_ignores_workbook_ids_and_stays_in_the_callers_organization(
    client: TestClient, tenants, agent_sessions: sessionmaker[Session]
) -> None:
    a, b = tenants
    # The workbook reuses B's database IDs; none of them may be touched or linked.
    assert b.customer_id == 1
    upload = client.post(
        "/api/import",
        files={"file": ("crm.xlsx", imported_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert upload.status_code == 200, upload.text
    with agent_sessions() as db:
        imported = db.scalar(select(Customer).where(Customer.customer_name == "Imported Customer"))
        assert imported.organization_id == a.organization_id
        assert imported.customer_id not in {a.customer_id, b.customer_id}
        assert db.get(Customer, b.customer_id).customer_name == "bravo Customer"
        meeting = db.scalar(select(Meeting).where(Meeting.customer_id == imported.customer_id))
        assert meeting.organization_id == a.organization_id
        assert db.get(Contact, meeting.contact_id).customer_id == imported.customer_id
        for model in (Contact, SalesEnquiry, Meeting, Call, FollowUp):
            assert db.scalar(
                select(func.count()).select_from(model).where(model.organization_id == b.organization_id)
            ) == 1


def test_import_rejects_duplicate_workbook_ids(client: TestClient, tenants) -> None:
    workbook = imported_workbook(include_second_customer=True)
    from io import BytesIO
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(workbook))
    book["Customers"].cell(row=3, column=1, value=1)  # Second customer reuses ID 1.
    output = BytesIO()
    book.save(output)
    response = client.post(
        "/api/import",
        files={"file": ("crm.xlsx", output.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 422
    assert "duplicate customer_id" in response.json()["error"]["message"]


def test_calendar_sync_never_exports_foreign_activities(
    google: FakeGoogle, tenants, agent_sessions: sessionmaker[Session]  # noqa: F811
) -> None:
    a, b = tenants
    with agent_sessions() as db:
        db.add(GoogleCalendarConnection(
            user_id=a.user_id,
            encrypted_refresh_token=calendar_sync_service._encrypt("refresh"),
            encrypted_access_token=calendar_sync_service._encrypt("access"),
            access_token_expires_at=datetime(2099, 1, 1, tzinfo=timezone.utc),
            scope="calendar",
            reminder_minutes=10,
        ))
        db.commit()
        calendar_sync_service.sync_activity(db, "meeting", b.meeting_id, created_by_user_id=a.user_id)
        assert google.calls == []
        assert db.scalar(select(func.count()).select_from(CalendarEventLink)) == 0
        # Control: A's own activity does sync.
        calendar_sync_service.sync_activity(db, "meeting", a.meeting_id, created_by_user_id=a.user_id)
        assert [call[0] for call in google.calls] == ["insert"]
        assert "bravo" not in json.dumps(google.events)
